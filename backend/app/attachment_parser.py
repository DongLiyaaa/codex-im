"""Bounded document parser CLI. No macros, formulas, links or embedded programs execute."""
import csv
import hashlib
import io
import json
import os
from pathlib import Path, PurePosixPath
import resource
import sys
import zipfile


def bound(name, default):
    return int(os.getenv('ATTACHMENT_' + name, str(default)))


def fail(message):
    raise ValueError(message)


def check_zip(raw):
    from defusedxml import ElementTree
    archive = zipfile.ZipFile(io.BytesIO(raw))
    infos = archive.infolist()
    if len(infos) > bound('ZIP_ENTRIES', 2000):
        fail('压缩包条目过多')
    total, names = 0, set()
    for entry in infos:
        name = entry.filename
        if name in names or '\\' in name or PurePosixPath(name).is_absolute() or '..' in PurePosixPath(name).parts or (entry.external_attr >> 16) & 0o170000 == 0o120000:
            fail('压缩包含重复或非法路径')
        names.add(name)
        total += entry.file_size
        if entry.flag_bits & 1 or entry.file_size > bound('ZIP_MEMBER_BYTES', 20 * 1024 * 1024) or total > bound('ZIP_EXPANDED_BYTES', 64 * 1024 * 1024) or entry.file_size > max(1, entry.compress_size) * 200:
            fail('压缩包加密、展开体积或压缩比超过限制')
        if name.lower().endswith(('vbaproject.bin', '.exe', '.dll')):
            fail('暂不支持含宏或可执行程序的附件')
        if name.endswith(('.xml', '.rels')):
            ElementTree.fromstring(archive.read(entry))
    return archive, names


def parse(source, output, filename=''):
    raw = source.read_bytes()
    if len(raw) > bound('MAX_BYTES', 20 * 1024 * 1024) or not raw:
        fail('附件为空或超过大小限制')
    data = {'pages': [], 'sheets': [], 'images': [], 'warnings': []}
    if raw.startswith(b'\xd0\xcf\x11\xe0'):
        fail('暂不支持旧版DOC/XLS，请另存为DOCX/XLSX；隔离转换将在后续提供')
    if raw.startswith((b'\x89PNG\r\n\x1a\n', b'\xff\xd8\xff')) or (raw[:4] == b'RIFF' and raw[8:12] == b'WEBP'):
        from PIL import Image, ImageOps
        Image.MAX_IMAGE_PIXELS = bound('IMAGE_PIXELS', 25_000_000)
        import warnings
        warnings.simplefilter('error', Image.DecompressionBombWarning)
        with Image.open(io.BytesIO(raw)) as image:
            if image.format not in ('JPEG', 'PNG', 'WEBP') or image.width * image.height > Image.MAX_IMAGE_PIXELS:
                fail('图片格式或像素超过限制')
            image.verify()
        with Image.open(io.BytesIO(raw)) as image:
            safe = ImageOps.exif_transpose(image).convert('RGB')
            safe.thumbnail((4096, 4096))
            target = output.parent / 'sanitized.png'
            safe.save(target, format='PNG')
            target.chmod(0o660)
            sanitized = target.read_bytes()
            if len(sanitized) > bound('MAX_BYTES', 20 * 1024 * 1024):
                fail('图片规范化后体积超过限制')
            data.update(kind='image', mime='image/png', images=[{'name': 'sanitized.png', 'size': len(sanitized), 'sha256': hashlib.sha256(sanitized).hexdigest(), 'mime': 'image/png'}])
        return data
    if raw.startswith(b'%PDF-'):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw), strict=True)
        if reader.is_encrypted:
            fail('暂不支持加密PDF')
        if len(reader.pages) > bound('PDF_PAGES', 200):
            fail('PDF页数超过限制')
        total = 0
        for i, page in enumerate(reader.pages, 1):
            contents = page.get_contents()
            if contents and len(contents.get_data()) > bound('PDF_STREAM_BYTES', 16 * 1024 * 1024):
                fail('PDF页展开内容超过限制')
            text = page.extract_text() or ''
            total += len(text)
            if total > bound('TEXT_CHARS', 2_000_000):
                fail('PDF文字内容超过限制')
            if not text.strip():
                fail('PDF含无文字层页面；当前未启用OCR，请上传文字版PDF或将所需页面转为图片单独分析')
            data['pages'].append({'page': i, 'text': text})
        data.update(kind='pdf', mime='application/pdf')
        return data
    if raw.startswith(b'PK\x03\x04'):
        archive, names = check_zip(raw)
        if 'word/document.xml' in names and 'xl/workbook.xml' not in names:
            from docx import Document
            document = Document(io.BytesIO(raw))
            text = '\n'.join(p.text for p in document.paragraphs)
            for table in document.tables:
                text += '\n' + '\n'.join('\t'.join(c.text for c in row.cells) for row in table.rows)
            if len(text) > bound('TEXT_CHARS', 2_000_000):
                fail('DOCX文字内容超过限制')
            if not text.strip():
                fail('DOCX没有可读取文字，嵌入图片暂不提取')
            data.update(kind='docx', mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document', pages=[{'page': 1, 'text': text}])
            data['warnings'] = ['DOCX页码为逻辑文字块，不代表Word排版页；嵌入图片、页眉页脚和文本框暂不提取']
            return data
        if 'xl/workbook.xml' in names and 'word/document.xml' not in names:
            from openpyxl import load_workbook
            values = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
            formulas = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            count = 0
            try:
                if len(formulas.worksheets) > bound('SHEETS', 50):
                    fail('工作表数量超过限制')
                for sheet in formulas:
                    sheet.reset_dimensions()
                    values[sheet.title].reset_dimensions()
                    rows = []
                    for row, cached in zip(sheet.iter_rows(), values[sheet.title].iter_rows()):
                        count += len(row)
                        if count > bound('CELLS', 100000) or len(row) > 1000:
                            fail('表格单元格或列数超过限制')
                        cells = []
                        for cell, cache in zip(row, cached):
                            value = cell.value
                            if cell.data_type == 'f':
                                value = {'formula': str(value), 'cached': cache.value, 'executed': False}
                            elif value is not None and not isinstance(value, (str, int, float, bool)):
                                value = str(value)
                            cells.append(value)
                        rows.append(cells)
                    data['sheets'].append({'name': sheet.title, 'rows': rows})
            finally:
                values.close()
                formulas.close()
            data.update(kind='xlsx', mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            return data
        fail('压缩包不是受支持的DOCX/XLSX')
    # Text has no magic signature. Strictly decode and reject control/binary data.
    if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
        text = raw.decode('utf-16', errors='strict')
    else:
        try:
            text = raw.decode('utf-8-sig', errors='strict')
        except UnicodeDecodeError:
            text = raw.decode('gb18030', errors='strict')
    if any(ord(c) < 32 and c not in '\n\r\t' for c in text) or len(text) > bound('TEXT_CHARS', 2_000_000):
        fail('文件不是安全文本或文本超过限制')
    if not text.strip() or text.lstrip().lower().startswith(('<!doctype html', '<html', '<?xml', '<svg')):
        fail('附件格式不受支持')
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        rows = list(csv.reader(io.StringIO(text), dialect=dialect))
        is_csv = len(rows) > 1 and len(rows[0]) > 1 and len({len(row) for row in rows}) == 1
    except csv.Error:
        is_csv = False
    if is_csv:
        if sum(map(len, rows)) > bound('CELLS', 100000):
            fail('CSV单元格超过限制')
        data.update(kind='csv', mime='text/csv', sheets=[{'name': 'CSV', 'rows': rows}])
    else:
        data.update(kind='txt', mime='text/plain', pages=[{'page': 1, 'text': text}])
    return data


def main():
    source, output = map(Path, sys.argv[1:3])
    resource.setrlimit(resource.RLIMIT_CPU, (40, 40))
    resource.setrlimit(resource.RLIMIT_FSIZE, (bound('RESULT_BYTES', 8 * 1024 * 1024),) * 2)
    resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
    if sys.platform == 'linux':
        resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024,) * 2)
    try:
        value = parse(source, output)
        code = 0
    except Exception as exc:
        value = {'error': str(exc)[:280] if isinstance(exc, ValueError) else '附件格式损坏、编码无效或解析超出安全限制'}
        code = 1
    encoded = json.dumps(value, ensure_ascii=False, default=str).encode()
    if len(encoded) > bound('RESULT_BYTES', 8 * 1024 * 1024):
        encoded = json.dumps({'error': '附件解析结果超过限制'}).encode()
        code = 1
    output.write_bytes(encoded)
    output.chmod(0o660)
    raise SystemExit(code)


if __name__ == '__main__':
    main()
