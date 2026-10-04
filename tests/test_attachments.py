"""Attachment trust boundaries, parser limits, queue and capability integration."""
import io
import json
from pathlib import Path
import zipfile
import pytest
from fastapi import HTTPException
from sqlalchemy import select
from test_im_postgres import database
from app.models import Conversation, User, Run, uid
from app.attachment_models import Attachment, AttachmentJob, AttachmentArtifact
from app import attachments, attachment_parser, attachment_worker, attachment_bridge, attachment_download, service


def test_parser_formats_and_formula_cache(tmp_path):
    from PIL import Image
    from docx import Document
    from openpyxl import Workbook
    from pypdf import PdfWriter
    source, output = tmp_path / 'source', tmp_path / 'result.json'
    source.write_text('区域,金额\n华东,137\n华南,263\n')
    assert attachment_parser.parse(source, output)['sheets'][0]['rows'][1][1] == '137'
    source.write_text('测试文本内容 481')
    assert '481' in attachment_parser.parse(source, output)['pages'][0]['text']
    image = Image.new('RGB', (30, 20), 'red'); image.save(source, format='PNG')
    assert attachment_parser.parse(source, output)['images'][0]['sha256']
    document = Document(); document.add_paragraph('订单金额 317'); document.save(source)
    assert '317' in attachment_parser.parse(source, output)['pages'][0]['text']
    book = Workbook(); book.active.append(['amount', 19, '=B1*2']); book.save(source)
    cell = attachment_parser.parse(source, output)['sheets'][0]['rows'][0][2]
    assert cell == {'formula': '=B1*2', 'cached': None, 'executed': False}
    writer = PdfWriter(); writer.add_blank_page(width=100, height=100); writer.write(source)
    with pytest.raises(ValueError, match='OCR'):
        attachment_parser.parse(source, output)


@pytest.mark.parametrize('name,body', [('../evil.xml', b'<a/>'), ('safe.xml', b'<!DOCTYPE a [<!ENTITY x SYSTEM "file:///etc/passwd">]><a>&x;</a>'), ('word/vbaProject.bin', b'macro')])
def test_zip_rejects_paths_xxe_macros(name, body):
    raw = io.BytesIO()
    with zipfile.ZipFile(raw, 'w') as z:
        z.writestr(name, body)
    with pytest.raises(Exception):
        attachment_parser.check_zip(raw.getvalue())


def test_zip_bomb_and_image_limits(tmp_path, monkeypatch):
    raw = io.BytesIO()
    with zipfile.ZipFile(raw, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('bomb', '0' * 100000)
    with pytest.raises(ValueError):
        attachment_parser.check_zip(raw.getvalue())
    from PIL import Image
    source = tmp_path / 'source'; Image.new('RGB', (20, 20)).save(source, format='PNG')
    monkeypatch.setenv('ATTACHMENT_IMAGE_PIXELS', '100')
    with pytest.raises(Exception):
        attachment_parser.parse(source, tmp_path / 'result.json')


def seed(db):
    user = db.scalar(select(User))
    conversation = Conversation(title='attachment-test', owner_id=user.id)
    db.add(conversation); db.flush()
    item = Attachment(id=uid(), conversation_id=conversation.id, uploader_id=user.id, filename='sample.txt', status='received')
    db.add(item); db.flush(); db.add(AttachmentJob(attachment_id=item.id))
    return user, conversation, item


def test_waiting_claim_lease_and_capability(database, monkeypatch):
    monkeypatch.setenv('PLATFORM_BRIDGE_KEY', 'unit-test-bridge-key-at-least-32-characters')
    with database.begin() as db:
        user, conv, item = seed(db)
        result = service.enqueue_message(db, user, conv, '', [item.id])
        run_id, item_id = result['run']['id'], item.id
        assert result['run']['status'] == 'waiting_attachments'
        with pytest.raises(HTTPException):
            service.archive_conversation(db, user, conv.id)
    ticket = attachment_worker.lease(database)
    assert ticket and attachment_worker.lease(database) is None
    with database.begin() as db:
        item = db.get(Attachment, item_id); item.status = 'ready'
        db.add(AttachmentArtifact(attachment_id=item_id, manifest={'pages': [{'page': 1, 'text': 'invoice 317'}]}))
    attachment_worker.wake_runs(database)
    with database.begin() as db:
        run = db.get(Run, run_id); assert run.status == 'queued'; run.status = 'running'; db.flush()
        token = attachment_bridge.issue(run)
        assert attachment_bridge.verify(token, db).id == run.id
        data = attachment_bridge.operate(db, run, 'read_document', {'attachment_id': item_id, 'limit': 5})
        assert data['text'] == 'invoi' and data['truncated']
        with pytest.raises(HTTPException): attachment_bridge.verify(token + 'a', db)
        with pytest.raises(HTTPException): attachment_bridge.verify(token, db, attachment_bridge.IMAGE_AUDIENCE)
        run.status = 'succeeded'; db.flush()
        with pytest.raises(HTTPException): attachment_bridge.verify(token, db)


def test_cross_user_and_conversation_claim(database):
    with database.begin() as db:
        user, conv, item = seed(db)
        other = User(email='other@example.invalid', name='other', role='member', org_id='org', team_id='team', password_hash='x')
        db.add(other); db.flush()
        with pytest.raises(HTTPException): attachments.authorize(db, item, actor=other)
        second = Conversation(title='second', owner_id=user.id); db.add(second); db.flush()
        with pytest.raises(HTTPException): service.enqueue_message(db, user, second, 'hi', [item.id])


@pytest.mark.parametrize('url', ['http://open.feishu.cn/file', 'https://127.0.0.1/a', 'https://open.feishu.cn.evil.test/a', 'https://user:pass@open.feishu.cn/a'])
def test_ssrf_url_rejection(url):
    with pytest.raises(RuntimeError): attachment_download.request(url, allowed=('open.feishu.cn',))


@pytest.mark.parametrize('address', ['127.0.0.1', '10.0.0.1', '169.254.169.254', '192.0.2.1', '::1'])
def test_ssrf_private_dns(monkeypatch, address):
    monkeypatch.setattr(attachment_download.socket, 'getaddrinfo', lambda *a, **k: [(2, 1, 6, '', (address, 443))])
    with pytest.raises(RuntimeError, match='受限网络'):
        attachment_download.request('https://open.feishu.cn/a', allowed=('open.feishu.cn',))


@pytest.mark.parametrize('answer', ['8.8.8.8', '127.0.0.1', '198.18.0.1'])
def test_fake_ip_resolution_pins_only_public_official_host(monkeypatch, answer):
    calls = []
    monkeypatch.setattr(attachment_download.socket, 'getaddrinfo', lambda *a, **k: [(2, 1, 6, '', ('198.18.0.8', 443))])
    class Connection:
        def __init__(self, host, address): calls.append((host, address))
        def request(self, method, path, **kwargs): calls.append((method, path))
        def getresponse(self): return self
        status = 200
        def read(self, n): return json.dumps({'Status': 0, 'Answer': [{'type': 1, 'data': answer}]}).encode()
        def close(self): pass
    monkeypatch.setattr(attachment_download, 'PinnedHTTPS', Connection)
    if answer == '8.8.8.8':
        assert attachment_download.resolve_addresses('open.feishu.cn') == [answer]
    else:
        with pytest.raises(RuntimeError, match='受限网络'):
            attachment_download.resolve_addresses('open.feishu.cn')
    assert calls[0] == ('cloudflare-dns.com', '1.1.1.1')
    assert calls[1] == ('GET', '/dns-query?name=open.feishu.cn&type=A')
    calls.clear()
    with pytest.raises(RuntimeError): attachment_download.resolve_addresses('other.feishu.cn')
    assert calls == []
    if answer == '8.8.8.8':
        assert attachment_download.resolve_addresses('accounts.feishu.cn') == [answer]
        assert ('GET', '/dns-query?name=accounts.feishu.cn&type=A') in calls


@pytest.mark.parametrize('status,retryable', [(403, False), (404, False), (429, True), (503, True)])
def test_download_http_error_is_redacted_and_bounded(monkeypatch, status, retryable):
    monkeypatch.setattr(attachment_download, 'resolve_addresses', lambda host: ['8.8.8.8'])
    class Connection:
        def __init__(self, *args): pass
        def request(self, *args, **kwargs): pass
        def getresponse(self): return self
        def read(self, n): return b'{"code":99991672,"msg":"SECRET URL AND TOKEN"}'
        def close(self): pass
    Connection.status = status
    monkeypatch.setattr(attachment_download, 'PinnedHTTPS', Connection)
    with pytest.raises(attachment_download.HTTPDownloadError) as error:
        attachment_download.request('https://open.feishu.cn/a', allowed=('open.feishu.cn',))
    assert error.value.retryable == retryable
    assert str(error.value) == f'ATTACHMENT_HTTP_{status}_CODE_99991672'


def test_feishu_image_only_fake_dns_worker_pipeline(database, tmp_path, monkeypatch):
    from PIL import Image
    from app import im, attachment_ingress
    from app.models import IMEvent
    monkeypatch.setenv('ATTACHMENT_ROOT', str(tmp_path))
    image = io.BytesIO()
    Image.new('RGB', (16, 16), 'blue').save(image, format='JPEG')
    monkeypatch.setattr(im, 'access_token', lambda *args: 'mock-token')
    monkeypatch.setattr(attachment_download.socket, 'getaddrinfo', lambda *a, **k: [(2, 1, 6, '', ('198.18.0.8', 443))])
    connections = []
    class Connection:
        status = 200
        def __init__(self, host, address):
            connections.append((host, address))
            self.data = io.BytesIO(json.dumps({'Status': 0, 'Answer': [{'type': 1, 'data': '8.8.8.8'}]}).encode() if host == 'cloudflare-dns.com' else image.getvalue())
        def request(self, method, path, **kwargs):
            if path.startswith('/open-apis/'):
                assert path == '/open-apis/im/v1/messages/image-event/resources/image-key?type=image'
        def getresponse(self): return self
        def getheader(self, name, default=None): return default
        def read(self, n): return self.data.read(n)
        def close(self): pass
    monkeypatch.setattr(attachment_download, 'PinnedHTTPS', Connection)
    text, refs = attachment_ingress.feishu({'message_type': 'image', 'message_id': 'image-event', 'content': json.dumps({'image_key': 'image-key'})})
    assert text == ''
    with database.begin() as db:
        im._enqueue(db, 'feishu', 'image-event', 'sender', 'chat', text, True, attachment_refs=refs)
        item = db.scalar(select(Attachment)); rid, aid = item.run_id, item.id
        assert db.get(Run, rid).status == 'waiting_attachments'
    attachment_worker.process(attachment_worker.lease(database), database)
    with database() as db:
        assert db.get(Attachment, aid).status == 'ready'
        assert db.get(Attachment, aid).mime == 'image/png'
        assert db.get(AttachmentJob, aid).state == 'done'
        assert db.get(Run, rid).status == 'queued'
        assert db.scalar(select(IMEvent)).delivered_at is None
    assert connections == [('cloudflare-dns.com', '1.1.1.1'), ('open.feishu.cn', '8.8.8.8')]


def test_upload_http_private_and_multipart(database, tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.db import get_db
    from app.security import current_user
    monkeypatch.setenv('ATTACHMENT_ROOT', str(tmp_path))
    with database.begin() as db:
        user, conv, _ = seed(db)
        actor, cid = user, conv.id
    app = FastAPI(); app.include_router(attachments.router)
    def session():
        with database.begin() as db:
            yield db
    app.dependency_overrides[get_db] = session
    app.dependency_overrides[current_user] = lambda: actor
    client = TestClient(app)
    # current_user normally shares the DB session; use a merged instance in this harness.
    def current():
        with database() as db:
            return db.get(User, actor.id)
    app.dependency_overrides[current_user] = current
    response = client.post(f'/api/conversations/{cid}/attachments', files={'file': ('note.txt', b'hello 619', 'text/plain')})
    assert response.status_code == 201, response.text
    aid = response.json()['id']
    assert client.get(f'/api/conversations/{cid}/attachments').status_code == 200
    assert client.delete(f'/api/conversations/{cid}/attachments/{aid}').status_code == 200


def test_worker_failed_parse_stops_run(database, tmp_path, monkeypatch):
    monkeypatch.setenv('ATTACHMENT_ROOT', str(tmp_path))
    with database.begin() as db:
        user, conv, item = seed(db)
        result = service.enqueue_message(db, user, conv, '', [item.id])
        rid = result['run']['id']
    def reject(_): raise RuntimeError('扫描PDF需要OCR')
    monkeypatch.setattr(attachment_worker, 'parse', reject)
    attachment_worker.process(attachment_worker.lease(database), database)
    with database() as db:
        assert db.get(Run, rid).status == 'failed'
        assert db.get(AttachmentJob, item.id).state == 'failed'


def test_im_attachment_dedup_unknown_and_authorized(database):
    from app.im import _enqueue
    refs = [{'filename': 'report.txt', 'reference': {'message_id': 'im-1', 'key': 'file-1', 'type': 'file'}}]
    with database.begin() as db:
        assert _enqueue(db, 'feishu', 'unknown-file', 'unknown', 'chat', '', True, attachment_refs=refs)['pending']
        assert list(db.scalars(select(Attachment))) == []
    with database.begin() as db:
        assert _enqueue(db, 'feishu', 'im-1', 'sender', 'chat', '', True, attachment_refs=refs)['ok']
    with database.begin() as db:
        assert _enqueue(db, 'feishu', 'im-1', 'sender', 'chat', '', True, attachment_refs=refs)['duplicate']
        rows = list(db.scalars(select(Attachment)))
        assert len(rows) == 1 and rows[0].run_id and rows[0].encrypted_reference
        assert db.get(Run, rows[0].run_id).status == 'waiting_attachments'


def test_cross_run_and_inactive_rejected(database, monkeypatch):
    monkeypatch.setenv('PLATFORM_BRIDGE_KEY', 'unit-test-bridge-key-at-least-32-characters')
    with database.begin() as db:
        user, conv, item = seed(db)
        result = service.enqueue_message(db, user, conv, '', [item.id])
        run = db.get(Run, result['run']['id']); run.status = 'running'; item.status = 'ready'; db.flush()
        other = Run(user_id=user.id, conversation_id=conv.id, message_id=run.message_id, status='running'); db.add(other); db.flush()
        with pytest.raises(HTTPException): attachment_bridge.bound(db, other, item.id)
        token = attachment_bridge.issue(run)
        user.active = False; db.flush()
        with pytest.raises(HTTPException): attachment_bridge.verify(token, db)


def test_subprocess_parser(tmp_path, monkeypatch):
    from app import attachment_storage
    monkeypatch.setenv('ATTACHMENT_ROOT', str(tmp_path))
    identifier = uid()
    attachment_storage.save_stream(identifier, [b'safe text 731'])
    result = attachment_worker.parse(identifier)
    assert '731' in result['pages'][0]['text']
