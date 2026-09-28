"""Synthetic-only real Codex attachment verification in an isolated PostgreSQL schema."""
import asyncio
import importlib.util
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import time
import pytest
import uvicorn
from fastapi import FastAPI
from sqlalchemy import select
ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'backend'), str(ROOT / 'tests')]
from test_im_postgres import database
from app import attachment_bridge, attachment_storage, attachment_worker, service
from app.attachment_models import Attachment, AttachmentArtifact
from app.models import User, Conversation, Message, Run, Audit, uid


def main():
    for line in (ROOT / '.env').read_text().splitlines():
        if line and not line.startswith('#') and '=' in line:
            k, v = line.split('=', 1)
            if k != 'DATABASE_URL': os.environ[k] = v
    os.environ['ATTACHMENT_BRIDGE_URL'] = 'http://127.0.0.1:18210/internal/attachment-mcp'
    os.environ['CODEX_OAUTH_AUTH_FILE'] = str(ROOT / '.runtime/codex-oauth/auth.json')
    os.environ['PATH'] = str(ROOT / '.runtime/codex-cli/node_modules/.bin') + os.pathsep + os.environ['PATH']
    patch = pytest.MonkeyPatch(); fixture = database.__wrapped__(patch); factory = next(fixture)
    app = FastAPI(); app.include_router(attachment_bridge.router)
    def session():
        with factory.begin() as db: yield db
    app.dependency_overrides[attachment_bridge.get_db] = session
    server = uvicorn.Server(uvicorn.Config(app, host='127.0.0.1', port=18210, log_level='error', access_log=False))
    thread = threading.Thread(target=server.run, daemon=True); thread.start()
    try:
        with tempfile.TemporaryDirectory(prefix='hub-attachment-e2e-') as temporary:
            os.environ['ATTACHMENT_ROOT'] = str(Path(temporary).resolve())
            with factory.begin() as db:
                user = db.scalar(select(User))
                conv = Conversation(owner_id=user.id, title='synthetic attachment E2E'); db.add(conv); db.flush()
                message = Message(conversation_id=conv.id, role='user', content='分析本次附件'); db.add(message); db.flush()
                run = Run(user_id=user.id, conversation_id=conv.id, message_id=message.id, status='running'); db.add(run); db.flush()
                files = ['visual.png', 'table.csv', 'notes.txt', 'report.docx', 'sales.xlsx', 'invoice.pdf']
                for filename in files:
                    identifier = uid(); source = attachment_storage.path(identifier)
                    if filename.endswith('.png'):
                        from PIL import Image, ImageDraw, ImageFont
                        image = Image.new('RGB', (640, 180), 'white')
                        font = ImageFont.truetype('/System/Library/Fonts/Helvetica.ttc', 42)
                        ImageDraw.Draw(image).text((20, 50), 'VISUAL-CHECK 8426', fill='black', font=font); image.save(source, format='PNG')
                    elif filename.endswith('.csv'): source.write_text('region,amount\nEast,137\nSouth,263\n')
                    elif filename.endswith('.txt'): source.write_text('Reference code: TEXT-619.\n')
                    elif filename.endswith('.docx'):
                        from docx import Document
                        document = Document(); document.add_paragraph('Document invoice total: 317 USD.'); document.save(source)
                    elif filename.endswith('.xlsx'):
                        from openpyxl import Workbook
                        book = Workbook(); book.active.title = 'Sales'; book.active.append(['Product', 'Units']); book.active.append(['Widget', 29]); book.save(source)
                    else:
                        from pypdf import PdfWriter
                        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
                        writer = PdfWriter(); page = writer.add_blank_page(width=500, height=200)
                        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
                        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
                        stream = DecodedStreamObject(); stream.set_data(b'BT /F1 20 Tf 20 100 Td (PDF total: 528 USD) Tj ET')
                        page[NameObject('/Contents')] = writer._add_object(stream); writer.write(source)
                    manifest = attachment_worker.parse(identifier)
                    db.add(Attachment(id=identifier, conversation_id=conv.id, uploader_id=user.id, message_id=message.id, run_id=run.id, filename=filename, mime=manifest['mime'], size=source.stat().st_size, status='ready'))
                    db.flush(); db.add(AttachmentArtifact(attachment_id=identifier, manifest=manifest))
                db.flush()
                payload = service.build_payload(db, run)
                payload['attachment_capability'] = attachment_bridge.issue(run)
                payload['image_capability'] = attachment_bridge.issue(run, attachment_bridge.IMAGE_AUDIENCE)
                payload['images'] = [{'attachment_id': a.id, 'filename': a.filename, 'size': image['size'], 'sha256': image['sha256']} for a in db.scalars(select(Attachment)) for image in db.get(AttachmentArtifact, a.id).manifest.get('images', [])]
                payload['prompt'] += '\n请真实看图并逐个读取六个附件：报告图片内验证码、CSV金额合计、TXT参考码、DOCX发票金额、XLSX的Widget数量、PDF金额。每项引用文件名及页/表/范围。不得猜测未读内容。'
            for _ in range(50):
                if server.started: break
                time.sleep(.1)
            spec = importlib.util.spec_from_file_location('attachment_real_runner', ROOT / 'runner/main.py')
            runner = importlib.util.module_from_spec(spec); sys.modules[spec.name] = runner; spec.loader.exec_module(runner)
            result = asyncio.run(runner.execute(runner.Execute(**payload), None))
            print(result)
            for expected in ['8426', '400', '619', '317', '29', '528']:
                assert expected in result, 'Missing synthetic result: ' + expected
            with factory() as db:
                calls = [a.details for a in db.scalars(select(Audit).where(Audit.action == 'attachment.tool_read'))]
                print('Verified attachment tool calls:', json.dumps(calls))
                assert any(c['tool'] == 'read_document' for c in calls)
                assert any(c['tool'] == 'read_sheet_range' for c in calls)
            print('REAL_CODEX_ATTACHMENTS_PASS; synthetic files only; isolated PostgreSQL; no IM delivery')
    finally:
        server.should_exit = True; thread.join(5)
        try: next(fixture)
        except StopIteration: pass
        patch.undo()


if __name__ == '__main__': main()
